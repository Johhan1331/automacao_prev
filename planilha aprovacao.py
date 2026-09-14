from pathlib import Path
from docx import Document
from datetime import date
from openpyxl import Workbook
from abrir_pasta import pasta_manha
from openpyxl import styles
# IMPORTAÇÕES NECESSÁRIAS ANTES DO CÓDIGO

data_planilha = date.today()
planilha_gestao = Workbook()
aba = planilha_gestao.active
aba.title = "Aprovação"

# EDIÇÃO DOS TEXTOS PADRÕES

aba["A1"] = "Empresa"
aba["B1"] = "Data"
aba["C1"] = "Lotes/Documentos"
aba["D1"] = "Grupo de Pagamento"
aba["D2"] = "PORTABILIDADE RG"
aba["D3"] = "RESGATE RG"
aba["D4"] = "DEVOLUÇÃO RG"
aba["D5"] = "PORTABILIDADE IS"
aba["D6"] = "PORTABILIDADE IS"
aba["D7"] = "RESGATE IS"
aba["D8"] = "RESGATE IS"
aba["D9"] = "DEVOLUÇÃO IS"
aba["E1"] = "Aprovador 1 - Aline"
aba["F1"] = "Aprovador 2 - Elaine"

for colunasRG in range (2, 5):
	aba[f"A{colunasRG}"] = 1013
	
for colunasIS in range (5, 10):
	aba[f"A{colunasIS}"] = 1010

for datas in range (2, 10):
	aba[f"B{datas}"] = data_planilha
	aba[f"B{datas}"].number_format = "DD/MM/YYYY"

# EDIÇÃO DO DIMENSIONAMENTO DAS CÉLULAS E CENTRALIZAÇÃO	
for formatWidth in ["A", "B"]:
    aba.column_dimensions[formatWidth].width = 15

aba.column_dimensions["C"].width = 45

for formatWidth in ["D", "E", "F"]:
    aba.column_dimensions[formatWidth].width = 20
    
    
for formatAlign in aba.iter_rows(min_row=1, max_row=15, min_col=1, max_col=6):
	for celulas in formatAlign:
		celulas.alignment = styles.Alignment (
			horizontal = "center",
			vertical = "center"
        )

# EDIÇÃO DAS BORDAS DAS CÉLULAS
bordas = styles.Side (
	style = "thin",
	color = "000000")

for borderAlign in aba.iter_rows(min_row=1, max_row=9, min_col=1, max_col=6):
	for celulas in borderAlign:
            celulas.border = styles.Border(
            left=bordas,
            right=bordas,
            top=bordas,
            bottom=bordas
        ) 

# EDIÇÃO DAS CORES DAS FONTES E PREENCHIMENTO DA CÉLULA(FILLHEADER)
for fonteRG in aba.iter_rows(min_row=2, max_row=4, min_col=1, max_col=6):
    for celulas in fonteRG:
        celulas.font = styles.Font(
            color="00B050"
        )

for fonteIS in aba.iter_rows(min_row=5, max_row=10, min_col=1, max_col=6):
    for celulas in fonteIS:
        celulas.font = styles.Font(
            color="203764"
        )
# FONTE NEGRITO NOS TÍTULOS
for fontBold in aba.iter_rows(
     min_row=1, 
     max_row=1, 
     min_col=1, 
     max_col=6):
    for celulas in fontBold:
          celulas.font = styles.Font (
                bold=True
          )

#PREENCHIMENTO CÉLULAS
for fillHeader in aba.iter_rows(min_row=1, max_row=1, min_col=1, max_col=6):
    for celulas in fillHeader:
        celulas.fill = styles.PatternFill(
            fill_type="solid",
            fgColor="D9E1F2"
        )


# CONFIGURAÇÃO DA ABA VALORES - CÓPIA TOTAL DA ABA APROVAÇÃO
aba_valores = planilha_gestao.copy_worksheet(aba)
aba_valores.title = "Valores"

# CONFIGURAÇÃO INDIVIDUAL DA ABA VALORES
aba_valores["E1"] = "VALOR"
aba_valores["D10"] = "TOTAL"
aba_valores["E10"] = "=SUM(E2:E9)"
aba_valores.delete_cols(6)

# FORMATAÇÃO DAS CÉLULAS
for total in range(2, 11):
    aba_valores[f"E{total}"].number_format = '_-[$R$-pt-BR]* #,##0.00_-;_-[$R$-pt-BR]* -#,##0.00_-;_-[$R$-pt-BR]* "-"??_-;_-@_-'

    for totalBold in aba_valores.iter_rows(
    min_row=10,
    max_row=10,
    min_col=4,
    max_col=5
):
        for celulas in totalBold:
            celulas.font = styles.Font(
            bold=True
        )
            
# PREENCHIMENTO DA CÉLULA
for totalFill in aba_valores.iter_rows(min_row=10, max_row=10, min_col=4, max_col=5):
    for celulas in totalFill:
        celulas.fill = styles.PatternFill(
            fill_type="solid",
            fgColor="D9E1F2"
        )

# CONFIGURAÇÃO DAS BORDAS
borda_total = styles.Side(
style="thin",
color="000000"
    )

for totalBorder in aba_valores.iter_rows(
    min_row=10,
    max_row=10,
    min_col=4,
    max_col=5
):
    for celulas in totalBorder:
        celulas.border = styles.Border(
            left=borda_total,
            right=borda_total,
            top=borda_total,
            bottom=borda_total
        )

planilha_gestao.save(pasta_manha / "Planilha de aprovação_Manhã.xlsx")