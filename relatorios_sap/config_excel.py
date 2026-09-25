from pathlib import Path
from datetime import date, datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import win32com.client as win32

# ============================================================
# CONFIGURAÇÃO DE PASTA
# ============================================================
PASTA_RAIZ = Path(r"X:\Comum-PagamentosDiarios")
MESES = [
    "01 - JANEIRO",
    "02 - FEVEREIRO",
    "03 - MARÇO",
    "04 - ABRIL",
    "05 - MAIO",
    "06 - JUNHO",
    "07 - JULHO",
    "08 - AGOSTO",
    "09 - SETEMBRO",
    "10 - OUTUBRO",
    "11 - NOVEMBRO",
    "12 - DEZEMBRO"
]

# ============================================================
# DATA DA PASTA
# ============================================================
# DATA TEMPORÁRIA PARA TESTES
data_atual = date(2026, 9, 26)
pasta_mes = MESES[data_atual.month - 1]
data_formatada = data_atual.strftime("%d.%m.%Y")

pasta_manha = (
    PASTA_RAIZ
    / str(data_atual.year)
    / pasta_mes
    / data_formatada
    / "MANHÃ"
)
# ============================================================
# HORÁRIO DA EXTRAÇÃO
# ============================================================

if datetime.now().hour < 10:
    HORARIO_RELATORIO = "8H"
else:
    HORARIO_RELATORIO = "11H"
    
# IDENTIFICAÇÃO DO EXCEL NA PASTA "MANHA"
arquivos_excel = [
    arquivo
    for arquivo in pasta_manha.glob("EXPORT_*.xlsx")
]

print(f"Pasta analisada:")
print(pasta_manha)
print()
print("Arquivos encontrados:")

for arquivo in arquivos_excel:
    print(arquivo.name)

arquivo_1010 = None
arquivo_1013 = None

for arquivo in arquivos_excel:
    planilha = load_workbook(
        arquivo,
        read_only=True,
        data_only=True
    ).active

    empresa = str(planilha["C4"].value).strip()

    print(f"{arquivo.name} → Empresa: {empresa}")

    if empresa == "1010":
        arquivo_1010 = arquivo

    elif empresa == "1013":
        arquivo_1013 = arquivo

print()
print("Arquivo 1010:", arquivo_1010.name if arquivo_1010 else "Não encontrado")
print("Arquivo 1013:", arquivo_1013.name if arquivo_1013 else "Não encontrado")

print()

planilha_1010 = load_workbook(
    arquivo_1010,
    read_only=True,
    data_only=True
).active

planilha_1013 = load_workbook(
    arquivo_1013,
    read_only=True,
    data_only=True
).active

print("Estrutura do arquivo 1010:")
print(f"Linhas: {planilha_1010.max_row}")
print(f"Colunas: {planilha_1010.max_column}")

print()

print("Estrutura do arquivo 1013:")
print(f"Linhas: {planilha_1013.max_row}")
print(f"Colunas: {planilha_1013.max_column}")    
# ============================================================
# CONFIGURAÇÃO DO EXCEL FINAL
# ============================================================

COLUNAS_FINAIS = [
    "Status",
    "Empresa",
    "Parceiro de Negócios",
    "Conta Contrato",
    "Objeto de Seguros",
    "Documento",
    "CPF-CNPJ",
    "Nome",
    "Tipo de Documento",
    "Origem",
    "Data do Documento",
    "Data de Lançamento",
    "Data de Vencimento",
    "Forma de Pagamento",
    "N° Lote de Pagamento",
    "Valor",
    "Banco de Pagamento",
    "Usuário Criação",
    "Objeto de bloqueio",
    "Categoria de bloqueio",
    "Processo de bloqueio",
    "Motivo do bloqueio",
    "Número do Voucher",
    "Banco Empresa",
    "Mensagem"
]

# ============================================================
# CRIA RELATORIO SEG
# ============================================================
novo_arquivo = Workbook()
nova_planilha = novo_arquivo.active
nova_planilha.title = "RESUMO"

# CABEÇALHO
for coluna, titulo in enumerate(COLUNAS_FINAIS, start=1):
    nova_planilha.cell(
        row=1,
        column=coluna,
        value=titulo
    )

# COPIAR DADOS DO EXCEL BRUTO
linha_destino = 2

for linha in planilha_1010.iter_rows(
    min_row=4,
    min_col=2,
    max_col=26,
    values_only=True
):
    for coluna_destino, valor in enumerate(linha, start=1):

        coluna_origem = coluna_destino + 1

# TIRA O NEGATIVO DA COLUNA VALOR
        if coluna_origem == 17 and isinstance(valor, str):
            valor = valor.replace(".", "").replace(",", ".")
            valor = abs(float(valor))

# REMOVE ESPAÇOS COLUNA EMPRESA
        if coluna_origem == 25 and isinstance(valor, str):
            valor = valor.strip()

        nova_planilha.cell(
            row=linha_destino,
            column=coluna_destino,
            value=valor
        )

    linha_destino += 1
    
# FORMATAR COLUNA VALOR COMO CONTABIL    
for linha in range(2, linha_destino):
    nova_planilha.cell(
        row=linha,
        column=16
    ).number_format = '_-[$R$-pt-BR]* #,##0.00_-;_-[$R$-pt-BR]* -#,##0.00_-;_-[$R$-pt-BR]* "-"??_-;_-@_-'
    
# ATIVAÇÃO DE FILTROS
nova_planilha.auto_filter.ref = f"A1:Y{linha_destino - 1}"

# AJUSTAR FONTE
for linha in nova_planilha.iter_rows():
    for celula in linha:
        celula.font = Font(
            name="Aptos Narrow",
            size=11
        )

# AJUSTAR COLUNAS
for coluna in range(3, 9):
    nova_planilha.column_dimensions[
        get_column_letter(coluna)
    ].width = 18

nova_planilha.column_dimensions["H"].width = 30
nova_planilha.column_dimensions["K"].width = 11
nova_planilha.column_dimensions["L"].width = 11
nova_planilha.column_dimensions["M"].width = 11
nova_planilha.column_dimensions["P"].width = 18

ARQUIVO_DESTINO = (
    pasta_manha
    / f"RELATORIO SEG {HORARIO_RELATORIO}.xlsx"
)

novo_arquivo.save(ARQUIVO_DESTINO)

print()
print("Arquivo 1010 criado:")
print(ARQUIVO_DESTINO)

# ============================================================
# CRIA RELATORIO RG
# ============================================================
novo_arquivo_1013 = Workbook()
nova_planilha_1013 = novo_arquivo_1013.active
nova_planilha_1013.title = "RESUMO"

# CABEÇALHO
for coluna, titulo in enumerate(COLUNAS_FINAIS, start=1):
    nova_planilha_1013.cell(
        row=1,
        column=coluna,
        value=titulo
    )

# COPIAR DADOS DO EXCEL BRUTO
linha_destino_1013 = 2

for linha in planilha_1013.iter_rows(
    min_row=4,
    min_col=2,
    max_col=26,
    values_only=True
):
    for coluna_destino, valor in enumerate(linha, start=1):

        coluna_origem = coluna_destino + 1

# TIRA O NEGATIVO DA COLUNA VALOR
        if coluna_origem == 17 and isinstance(valor, str):
            valor = valor.replace(".", "").replace(",", ".")
            valor = abs(float(valor))

# REMOVE ESPAÇOS COLUNA EMPRESA
        if coluna_origem == 25 and isinstance(valor, str):
            valor = valor.strip()

        nova_planilha_1013.cell(
            row=linha_destino_1013,
            column=coluna_destino,
            value=valor
        )

    linha_destino_1013 += 1

# FORMATAR COLUNA VALOR COMO CONTABIL   

for linha in range(2, linha_destino_1013):
    nova_planilha_1013.cell(
        row=linha,
        column=16
    ).number_format = '_-[$R$-pt-BR]* #,##0.00_-;_-[$R$-pt-BR]* -#,##0.00_-;_-[$R$-pt-BR]* "-"??_-;_-@_-'

# ATIVAÇÃO DE FILTROS

nova_planilha_1013.auto_filter.ref = (
    f"A1:Y{linha_destino_1013 - 1}"
)

# AJUSTAR FONTE

for linha in nova_planilha_1013.iter_rows():
    for celula in linha:
        celula.font = Font(
            name="Aptos Narrow",
            size=11
        )

# AJUSTAR COLUNAS

for coluna in range(3, 9):
    nova_planilha_1013.column_dimensions[
        get_column_letter(coluna)
    ].width = 18

nova_planilha_1013.column_dimensions["H"].width = 30
nova_planilha_1013.column_dimensions["K"].width = 11
nova_planilha_1013.column_dimensions["L"].width = 11
nova_planilha_1013.column_dimensions["M"].width = 11
nova_planilha_1013.column_dimensions["P"].width = 18

# SALVAR
ARQUIVO_DESTINO_1013 = (
    pasta_manha
    / f"RELATORIO RG {HORARIO_RELATORIO}.xlsx"
)

novo_arquivo_1013.save(ARQUIVO_DESTINO_1013)

print()
print("Arquivo 1013 criado:")
print(ARQUIVO_DESTINO_1013)

# ============================================================
# TABELA DINAMICA 1010
# ============================================================

def criar_tabela_dinamica(
    caminho_arquivo,
    nome_tabela,
    nome_aba="TABELA DINAMICA"
):

    excel = None
    arquivo_excel = None

    try:
        excel = win32.Dispatch("Excel.Application")
        excel.Visible = False

        arquivo_excel = excel.Workbooks.Open(
            str(caminho_arquivo)
        )

        planilha_excel = arquivo_excel.Worksheets("RESUMO")

        planilha_dinamica = arquivo_excel.Worksheets.Add()
        planilha_dinamica.Name = nome_aba

        ultima_linha = planilha_excel.Cells(
            planilha_excel.Rows.Count,
            16
        ).End(-4162).Row

        linha_dinamica = 2

        titulo = planilha_dinamica.Cells(
            linha_dinamica,
            1
        )

        titulo.Value = "RELATÓRIO SAP"
        titulo.Font.Bold = True
        titulo.Font.Color = 255

        origem = (
            f"'RESUMO'!R1C1:R{ultima_linha}C25"
        )

        cache = arquivo_excel.PivotCaches().Create(
            SourceType=1,
            SourceData=origem
        )

        destino = planilha_dinamica.Cells(
            linha_dinamica + 1,
            1
        )

        tabela_dinamica = cache.CreatePivotTable(
            TableDestination=destino,
            TableName=nome_tabela
        )

        campo_lote = tabela_dinamica.PivotFields(
            "N° Lote de Pagamento"
        )

        campo_lote.Orientation = 1

        campo_valor = tabela_dinamica.PivotFields(
            "Valor"
        )

        campo_valor.Orientation = 4
        campo_valor.Function = -4157

        campo_documento = tabela_dinamica.PivotFields(
            "Documento"
        )

        campo_documento.Orientation = 4
        campo_documento.Function = -4112

        arquivo_excel.Save()

        print(
            f"Tabela dinâmica criada: "
            f"{caminho_arquivo}"
        )

    finally:

        if arquivo_excel is not None:
            arquivo_excel.Close(
                SaveChanges=False
            )

        if excel is not None:
            excel.Quit()


# ============================================================
# CHAMADA DA TABELA DINAMICA 1010
# ============================================================

criar_tabela_dinamica(
    ARQUIVO_DESTINO,
    "TabelaDinamica1010",
    "TABELA DINAMICA GERAL"
)

# ============================================================
# SEPARA POR TIPO 1010
# ============================================================

planilha_separacao = nova_planilha

tipo_56 = []
tipo_62 = []
tipo_63 = []
tipo_65_79 = []

for linha in range(2, planilha_separacao.max_row + 1):

    tipo_documento = str(
        planilha_separacao.cell(
            row=linha,
            column=9
        ).value
    ).strip()

    if tipo_documento == "56":
        tipo_56.append(linha)

    elif tipo_documento == "62":
        tipo_62.append(linha)

    elif tipo_documento == "63":
        tipo_63.append(linha)

    elif tipo_documento in ("65", "79"):
        tipo_65_79.append(linha)

print()
print("SEPARAÇÃO 1010 POR TIPO DE DOCUMENTO")
print(f"Tipo 56: {len(tipo_56)} linhas")
print(f"Tipo 62: {len(tipo_62)} linhas")
print(f"Tipo 63: {len(tipo_63)} linhas")
print(f"Tipo 65 + 79: {len(tipo_65_79)} linhas")


# ============================================================
# CRIA EXCEL DOS TIPOS 1010
# ============================================================

def obter_lotes_ja_extraidos(descricao):

    lotes = set()

    pasta_tipo = pasta_manha / descricao

    if not pasta_tipo.exists():
        return lotes

    for arquivo in pasta_tipo.rglob("*.xlsx"):

        try:

            workbook = load_workbook(
                arquivo,
                read_only=True,
                data_only=True
            )

            planilha = workbook.active

            for linha in planilha.iter_rows(
                min_row=2,
                min_col=15,
                max_col=15,
                values_only=True
            ):

                lote = linha[0]

                if lote is not None:

                    lote = str(lote).strip()

                    if lote:
                        lotes.add(lote)

            workbook.close()

        except Exception:
            continue

    return lotes

# ============================================================
# LIMITAR NOME DOS LOTES
# ============================================================

LIMITE_CAMINHO = 240

def montar_nomes_com_lotes(
    lotes,
    pasta_tipo,
    descricao
):
    lotes_visiveis = []

    for lote in lotes:
        candidato_lotes = " ".join(
            lotes_visiveis + [lote]
        )

        nome_pasta = (
            f"LOTE {candidato_lotes} - {descricao}"
        )

        nome_arquivo = (
            f"LOTE {candidato_lotes} - {descricao}.xlsx"
        )

        caminho_candidato = (
            pasta_tipo
            / nome_pasta
            / nome_arquivo
        )

        if len(str(caminho_candidato)) > LIMITE_CAMINHO:
            break

        lotes_visiveis.append(lote)

    if len(lotes_visiveis) < len(lotes):
        lotes_nome = " ".join(lotes_visiveis) + "..."
    else:
        lotes_nome = " ".join(lotes_visiveis)

    nome_pasta = (
        f"LOTE {lotes_nome} - {descricao}"
    )

    nome_arquivo = (
        f"LOTE {lotes_nome} - {descricao}.xlsx"
    )

    return nome_pasta, nome_arquivo


def criar_arquivo_tipo(
    planilha_origem,
    linhas,
    descricao
):

    if not linhas:
        print(
            f"Nenhum registro encontrado para: "
            f"{descricao}"
        )
        return None

    lotes_ja_extraidos = obter_lotes_ja_extraidos(
        descricao
    )

    linhas_novas = []
    lotes = []

    for linha in linhas:

        lote = planilha_origem.cell(
            row=linha,
            column=15
        ).value

        lote = str(lote).strip()

        if lote in lotes_ja_extraidos:
            continue

        linhas_novas.append(linha)

        if lote not in lotes:
            lotes.append(lote)

    if not linhas_novas:

        print(
            f"Nenhum lote novo encontrado para: "
            f"{descricao}"
        )

        return None

# CRIAR PASTA DO TIPO
    pasta_tipo = pasta_manha / descricao
    pasta_tipo.mkdir(
    parents=True,
    exist_ok=True
)

# MONTAR NOMES DA PASTA E DO ARQUIVO
    nome_pasta, nome_arquivo = montar_nomes_com_lotes(
    lotes,
    pasta_tipo,
    descricao
)

# CRIAR PASTA DO RESULTADO
    pasta_resultado = pasta_tipo / nome_pasta
    pasta_resultado.mkdir(
    parents=True,
    exist_ok=True
)

# CAMINHO DO ARQUIVO
    caminho_arquivo = (
    pasta_resultado / nome_arquivo
)    

# CRIAR EXCEL

    arquivo = Workbook()

    planilha = arquivo.active
    planilha.title = "RESUMO"

    # COPIAR CABEÇALHO

    for coluna in range(
        1,
        planilha_origem.max_column + 1
    ):

        planilha.cell(
            row=1,
            column=coluna,
            value=planilha_origem.cell(
                row=1,
                column=coluna
            ).value
        )

    # COPIAR DOCUMENTOS NOVOS

    linha_destino = 2

    for linha_origem in linhas_novas:

        for coluna in range(
            1,
            planilha_origem.max_column + 1
        ):

            planilha.cell(
                row=linha_destino,
                column=coluna,
                value=planilha_origem.cell(
                    row=linha_origem,
                    column=coluna
                ).value
            )

        linha_destino += 1

    # FORMATAR COLUNA VALOR COMO CONTÁBIL

    for linha in range(
        2,
        linha_destino
    ):

        planilha.cell(
            row=linha,
            column=16
        ).number_format = (
            '_-[$R$-pt-BR]* #,##0.00_-;'
            '_-[$R$-pt-BR]* -#,##0.00_-;'
            '_-[$R$-pt-BR]* "-"??_-;'
            '_-@_-'
        )

    # ATIVAÇÃO DE FILTROS

    planilha.auto_filter.ref = (
        f"A1:Y{linha_destino - 1}"
    )

    # AJUSTAR FONTE

    for linha in planilha.iter_rows():

        for celula in linha:

            celula.font = Font(
                name="Aptos Narrow",
                size=11
            )

    # AJUSTAR COLUNAS

    for coluna in range(3, 9):

        planilha.column_dimensions[
            get_column_letter(coluna)
        ].width = 18

    planilha.column_dimensions["H"].width = 30
    planilha.column_dimensions["K"].width = 11
    planilha.column_dimensions["L"].width = 11
    planilha.column_dimensions["M"].width = 11
    planilha.column_dimensions["P"].width = 18

    arquivo.save(caminho_arquivo)
    arquivo.close()

    print()
    print(f"{descricao}")
    print(
        f"Lotes encontrados: "
        f"{' '.join(lotes)}"
    )
    print(
        f"Documentos: {len(linhas_novas)}"
    )
    print("Arquivo criado:")
    print(caminho_arquivo)

    return caminho_arquivo


# ============================================================
# TABELA DINAMICA DOS TIPOS 1010
# ============================================================

arquivo_portabilidade_is = criar_arquivo_tipo(
    planilha_separacao,
    tipo_56,
    "PORTABILIDADE IS"
)

if arquivo_portabilidade_is:

    criar_tabela_dinamica(
        arquivo_portabilidade_is,
        "TabelaDinamicaPortabilidadeIS"
    )


arquivo_resgate_is = criar_arquivo_tipo(
    planilha_separacao,
    tipo_62,
    "RESGATE IS"
)

if arquivo_resgate_is:

    criar_tabela_dinamica(
        arquivo_resgate_is,
        "TabelaDinamicaResgateIS"
    )


arquivo_renda_is = criar_arquivo_tipo(
    planilha_separacao,
    tipo_63,
    "RENDA IS"
)

if arquivo_renda_is:

    criar_tabela_dinamica(
        arquivo_renda_is,
        "TabelaDinamicaRendaIS"
    )


arquivo_devolucao_is = criar_arquivo_tipo(
    planilha_separacao,
    tipo_65_79,
    "DEVOLUÇÃO IS"
)

if arquivo_devolucao_is:

    criar_tabela_dinamica(
        arquivo_devolucao_is,
        "TabelaDinamicaDevolucaoIS"
    )


# ============================================================
# TABELA DINAMICA 1013
# ============================================================

criar_tabela_dinamica(
    ARQUIVO_DESTINO_1013,
    "TabelaDinamica1013",
    "TABELA DINAMICA GERAL"
)


# ============================================================
# SEPARA POR TIPO 1013
# ============================================================

planilha_separacao_1013 = nova_planilha_1013

tipo_56_1013 = []
tipo_62_1013 = []
tipo_63_1013 = []
tipo_65_1013 = []

for linha in range(
    2,
    planilha_separacao_1013.max_row + 1
):

    tipo_documento = str(
        planilha_separacao_1013.cell(
            row=linha,
            column=9
        ).value
    ).strip()

    if tipo_documento == "56":
        tipo_56_1013.append(linha)

    elif tipo_documento == "62":
        tipo_62_1013.append(linha)

    elif tipo_documento == "63":
        tipo_63_1013.append(linha)

    elif tipo_documento == "65":
        tipo_65_1013.append(linha)

print()
print("SEPARAÇÃO 1013 POR TIPO DE DOCUMENTO")
print(
    f"Tipo 56: {len(tipo_56_1013)} linhas"
)
print(
    f"Tipo 62: {len(tipo_62_1013)} linhas"
)
print(
    f"Tipo 63: {len(tipo_63_1013)} linhas"
)
print(
    f"Tipo 65: {len(tipo_65_1013)} linhas"
)


# ============================================================
# CRIA EXCEL DOS TIPOS 1013
# ============================================================

def criar_arquivo_tipo_1013(
    planilha_origem,
    linhas,
    descricao
):

    if not linhas:

        print(
            f"Nenhum registro encontrado para: "
            f"{descricao}"
        )

        return None

    lotes_ja_extraidos = obter_lotes_ja_extraidos(
        descricao
    )

    linhas_novas = []
    lotes = []

    for linha in linhas:

        lote = planilha_origem.cell(
            row=linha,
            column=15
        ).value

        lote = str(lote).strip()

        if lote in lotes_ja_extraidos:
            continue

        linhas_novas.append(linha)

        if lote not in lotes:
            lotes.append(lote)

    if not linhas_novas:

        print(
            f"Nenhum lote novo encontrado para: "
            f"{descricao}"
        )

        return None

    # MONTAR NOME DO ARQUIVO

    lotes_nome = " ".join(lotes)

    nome_arquivo = (
        f"LOTE {lotes_nome} - {descricao}.xlsx"
    )

    # CRIAR PASTA DO TIPO

    pasta_tipo = pasta_manha / descricao

    pasta_tipo.mkdir(
        parents=True,
        exist_ok=True
    )

    # CRIAR PASTA DO RESULTADO

    nome_pasta = (
        f"LOTE {lotes_nome} - {descricao}"
    )

    pasta_resultado = pasta_tipo / nome_pasta

    pasta_resultado.mkdir(
        parents=True,
        exist_ok=True
    )

    caminho_arquivo = (
        pasta_resultado / nome_arquivo
    )

    # CRIAR EXCEL

    arquivo = Workbook()

    planilha = arquivo.active
    planilha.title = "RESUMO"

    # COPIAR CABEÇALHO

    for coluna in range(
        1,
        planilha_origem.max_column + 1
    ):

        planilha.cell(
            row=1,
            column=coluna,
            value=planilha_origem.cell(
                row=1,
                column=coluna
            ).value
        )

    # COPIAR DOCUMENTOS NOVOS

    linha_destino = 2

    for linha_origem in linhas_novas:

        for coluna in range(
            1,
            planilha_origem.max_column + 1
        ):

            planilha.cell(
                row=linha_destino,
                column=coluna,
                value=planilha_origem.cell(
                    row=linha_origem,
                    column=coluna
                ).value
            )

        linha_destino += 1

    # FORMATAR COLUNA VALOR COMO CONTÁBIL

    for linha in range(
        2,
        linha_destino
    ):

        planilha.cell(
            row=linha,
            column=16
        ).number_format = (
            '_-[$R$-pt-BR]* #,##0.00_-;'
            '_-[$R$-pt-BR]* -#,##0.00_-;'
            '_-[$R$-pt-BR]* "-"??_-;'
            '_-@_-'
        )

    # ATIVAÇÃO DE FILTROS

    planilha.auto_filter.ref = (
        f"A1:Y{linha_destino - 1}"
    )

    # AJUSTAR FONTE

    for linha in planilha.iter_rows():

        for celula in linha:

            celula.font = Font(
                name="Aptos Narrow",
                size=11
            )

    # AJUSTAR COLUNAS

    for coluna in range(3, 9):

        planilha.column_dimensions[
            get_column_letter(coluna)
        ].width = 18

    planilha.column_dimensions["H"].width = 30
    planilha.column_dimensions["K"].width = 11
    planilha.column_dimensions["L"].width = 11
    planilha.column_dimensions["M"].width = 11
    planilha.column_dimensions["P"].width = 18

    arquivo.save(caminho_arquivo)

    print()
    print(f"{descricao}")
    print(
        f"Lotes encontrados: "
        f"{' '.join(lotes)}"
    )
    print(
        f"Documentos: {len(linhas_novas)}"
    )
    print("Arquivo criado:")
    print(caminho_arquivo)

    return caminho_arquivo


# ============================================================
# TABELA DINAMICA DOS TIPOS 1013
# ============================================================

arquivo_portabilidade_rg = criar_arquivo_tipo_1013(
    planilha_separacao_1013,
    tipo_56_1013,
    "PORTABILIDADE RG"
)

if arquivo_portabilidade_rg:

    criar_tabela_dinamica(
        arquivo_portabilidade_rg,
        "TabelaDinamicaPortabilidadeRG"
    )


arquivo_resgate_rg = criar_arquivo_tipo_1013(
    planilha_separacao_1013,
    tipo_62_1013,
    "RESGATE RG"
)

if arquivo_resgate_rg:

    criar_tabela_dinamica(
        arquivo_resgate_rg,
        "TabelaDinamicaResgateRG"
    )


arquivo_renda_rg = criar_arquivo_tipo_1013(
    planilha_separacao_1013,
    tipo_63_1013,
    "RENDA RG"
)

if arquivo_renda_rg:

    criar_tabela_dinamica(
        arquivo_renda_rg,
        "TabelaDinamicaRendaRG"
    )


arquivo_devolucao_rg = criar_arquivo_tipo_1013(
    planilha_separacao_1013,
    tipo_65_1013,
    "DEVOLUÇÃO RG"
)

if arquivo_devolucao_rg:

    criar_tabela_dinamica(
        arquivo_devolucao_rg,
        "TabelaDinamicaDevolucaoRG"
    )